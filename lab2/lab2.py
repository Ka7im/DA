# Работа 2. Агломеративная иерархическая кластеризация.
# Запуск всех пяти методов: python lab2.py --n-clusters 3

# Подключаем аргументы запуска и путь к Excel-файлу.
import argparse
from pathlib import Path

# Подключаем точные дроби, проверку чисел и вычисление медианы.
from fractions import Fraction
from math import isfinite
from numbers import Real
from statistics import median

# Подключаем чтение существующего Excel-файла.
from openpyxl import load_workbook


# ПУНКТ 1. КЛАСС АГЛОМЕРАТИВНОЙ ИЕРАРХИЧЕСКОЙ КЛАСТЕРИЗАЦИИ.
class AgglomerativeClustering:
    # Перечисляем методы в порядке задания.
    METHODS = {
        "single": "Односвязная кластеризация",
        "complete": "Полносвязная кластеризация",
        "average": "Среднее расстояние",
        "method1": "Метод 1 — полусумма минимума и максимума",
        "method2": "Метод 2 — медиана попарных расстояний",
    }

    # Принимаем матрицу, число оставшихся кластеров и выбранный метод.
    def __init__(self, distance_matrix, n_clusters=3,
                 method="single", labels=None):
        # Проверяем размер матрицы и допустимость параметров.
        n = len(distance_matrix)
        if n == 0 or any(len(row) != n for row in distance_matrix):
            raise ValueError("Матрица должна быть непустой и квадратной.")
        if type(n_clusters) is not int or not 1 <= n_clusters <= n:
            raise ValueError(f"n_clusters должно быть целым от 1 до {n}.")
        if method not in self.METHODS:
            raise ValueError("Неизвестный метод кластеризации.")

        # Сохраняем точные дроби, чтобы не искажать сравнения округлением.
        rows = []
        for row in distance_matrix:
            if any(isinstance(x, bool) or not isinstance(x, Real)
                   or not isfinite(x) or x < 0 for x in row):
                raise ValueError("Расстояния должны быть конечными и >= 0.")
            rows.append(tuple(Fraction(str(x)) for x in row))
        self.distances = tuple(rows)

        # Проверяем нулевую диагональ и симметрию исходных расстояний.
        for i in range(n):
            if self.distances[i][i] != 0:
                raise ValueError("Диагональ матрицы должна быть нулевой.")
            for j in range(i + 1, n):
                if self.distances[i][j] != self.distances[j][i]:
                    raise ValueError("Матрица должна быть симметричной.")

        # Подготавливаем уникальные подписи объектов и настройки.
        self.labels = tuple(labels if labels is not None else
                            (f"P{i + 1}" for i in range(n)))
        if (len(self.labels) != n
                or any(not isinstance(x, str) or not x for x in self.labels)
                or len(set(self.labels)) != n):
            raise ValueError("Нужна одна уникальная строковая метка на объект.")
        self.n_clusters = n_clusters
        self.method = method
        self.history = []
        self.clusters_ = {}

    # Вычисляем расстояние по ВСЕМ исходным парам двух кластеров.
    def cluster_distance(self, first, second):
        # Элементы first и second — индексы объектов от 0 до n - 1.
        values = [self.distances[i][j] for i in first for j in second]
        if not values:
            raise ValueError("Кластеры не должны быть пустыми.")

        # Односвязный метод: минимальное попарное расстояние.
        if self.method == "single":
            return min(values)

        # Полносвязный метод: максимальное попарное расстояние.
        if self.method == "complete":
            return max(values)

        # Среднее расстояние: сумма всех |A| * |B| значений / их число.
        if self.method == "average":
            return sum(values) / len(values)

        # Метод 1: полусумма минимального и максимального расстояний.
        if self.method == "method1":
            return (min(values) + max(values)) / 2

        # Метод 2: медиана исходных попарных расстояний с повторениями.
        # При чётном числе пар берётся полусумма двух средних значений.
        return median(values)

    # Формируем текущую матрицу в порядке возрастания номеров кластеров.
    def _cluster_matrix(self, clusters):
        ids = sorted(clusters)
        matrix = [[Fraction(0) for _ in ids] for _ in ids]
        for i, left in enumerate(ids):
            for j in range(i + 1, len(ids)):
                right = ids[j]
                distance = self.cluster_distance(clusters[left], clusters[right])
                matrix[i][j] = matrix[j][i] = distance
        return ids, matrix

    # Печатаем целые числа и точные дроби без потери точности.
    @staticmethod
    def print_matrix(matrix, labels):
        width = max(8, max(len(str(x)) for row in matrix for x in row) + 2,
                    max(len(x) for x in labels) + 2)
        print(" " * width + "".join(f"{x:>{width}}" for x in labels))
        for label, row in zip(labels, matrix):
            print(f"{label:>{width}}"
                  + "".join(f"{str(x):>{width}}" for x in row))

    # Преобразуем индексы объектов в понятный состав кластера.
    def _members_text(self, members):
        return "{" + ", ".join(self.labels[i] for i in members) + "}"

    # Выполняем объединения до заданного числа кластеров.
    def fit(self, verbose=True):
        # Каждый объект сначала образует свой кластер C1, ..., Cn.
        n = len(self.distances)
        clusters = {i + 1: (i,) for i in range(n)}
        self.history = []
        next_id = n + 1

        # Сначала выводим заданную на входе матрицу наблюдений.
        if verbose:
            print("Исходная матрица расстояний между наблюдениями:")
            self.print_matrix(self.distances, self.labels)
            print(f"Метод: {self.method} — {self.METHODS[self.method]}")
            print(f"Остановка при n_clusters = {self.n_clusters}")
            print("Начальные кластеры: " + "; ".join(
                f"C{key} = {self._members_text(value)}"
                for key, value in clusters.items()))

        # Строим матрицу расстояний между начальными кластерами.
        ids, matrix = self._cluster_matrix(clusters)
        while len(clusters) > self.n_clusters:
            # Ищем минимум только выше диагонали: кластер сам с собой не сливаем.
            # При равенстве расстояний выбираем меньшую пару номеров (a, b).
            distance, left, right = min(
                (matrix[i][j], ids[i], ids[j])
                for i in range(len(ids)) for j in range(i + 1, len(ids))
            )

            # Объединяем составы выбранных кластеров и присваиваем новый номер.
            first, second = clusters.pop(left), clusters.pop(right)
            clusters[next_id] = tuple(sorted(first + second))

            # Пересчитываем расстояния по исходным парам объектов.
            ids, matrix = self._cluster_matrix(clusters)
            step = len(self.history) + 1
            self.history.append({
                "step": step, "left": left, "right": right,
                "new_id": next_id, "distance": distance,
                "clusters": dict(clusters), "ids": tuple(ids),
                "matrix": tuple(tuple(row) for row in matrix),
            })

            # На каждом шаге выводим выбор, составы и обновлённую матрицу.
            if verbose:
                print(f"\nШаг {step}. C{left} + C{right} -> C{next_id}; "
                      f"расстояние = {distance}")
                print(f"Объединены {self._members_text(first)} и "
                      f"{self._members_text(second)}")
                print(f"Осталось кластеров: {len(clusters)}")
                for key in ids:
                    print(f"C{key} = {self._members_text(clusters[key])}")
                print("Матрица расстояний после объединения:")
                self.print_matrix(matrix, [f"C{key}" for key in ids])
            next_id += 1

        # Сохраняем результат и возвращаем кластеры с метками объектов.
        self.clusters_ = dict(clusters)
        result = [[self.labels[i] for i in clusters[key]] for key in ids]
        if verbose:
            print("\nИтог: " + "; ".join(self._members_text(clusters[key])
                                         for key in ids))
        return result


# ПУНКТ 2. МАТРИЦА ДЛЯ ВОСЬМИ НАБЛЮДЕНИЙ И ПРОВЕРКА ПЯТИ МЕТОДОВ.
# Матрица задана в lab2.xlsx, лист Матрица, диапазон B2:I9.
# Метки столбцов находятся в B1:I1, метки строк — в A2:A9.

# Читаем именно Excel-файл: готового списка исходных расстояний в коде нет.
def read_excel(filename):
    workbook = load_workbook(filename, read_only=True, data_only=True)
    try:
        # Получаем восемь меток столбцов и восемь строк с расстояниями.
        sheet = workbook["Матрица"]
        labels = [cell.value for cell in sheet["B1:I1"][0]]
        rows = list(sheet.iter_rows(min_row=2, max_row=9, max_col=9,
                                    values_only=True))

        # Проверяем согласованность подписей строк и столбцов.
        if [row[0] for row in rows] != labels:
            raise ValueError("Метки строк и столбцов Excel должны совпадать.")
        return [list(row[1:]) for row in rows], labels
    finally:
        # Закрываем Excel-файл даже при ошибке чтения.
        workbook.close()


# Выполняем демонстрацию при запуске этого файла как программы.
if __name__ == "__main__":
    # Задаём путь к матрице, число кластеров и один метод либо все методы.
    parser = argparse.ArgumentParser(description="Работа 2. Кластеризация")
    parser.add_argument("--file", type=Path,
                        default=Path(__file__).with_name("lab2.xlsx"))
    parser.add_argument("--n-clusters", type=int, default=3)
    parser.add_argument("--method", default="all",
                        choices=["all", *AgglomerativeClustering.METHODS])
    args = parser.parse_args()

    # Считываем общую исходную матрицу из Excel.
    distance_matrix, labels = read_excel(args.file)
    methods = (list(AgglomerativeClustering.METHODS) if args.method == "all"
               else [args.method])

    # Для каждого метода создаём новый объект с исходными одиночными кластерами.
    for method in methods:
        model = AgglomerativeClustering(distance_matrix, args.n_clusters,
                                        method, labels)
        model.fit()
        print()
