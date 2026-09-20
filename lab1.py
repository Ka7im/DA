# Работа 1. Матрица расстояний для восьми двумерных наблюдений.
# Запуск: python lab1.py

# Подключаем извлечение корня, проверку конечности числа и сравнение с допуском.
from math import sqrt, isfinite, isclose

# Подключаем работу с путями, чтобы найти Excel-файл рядом с программой.
from pathlib import Path

# Подключаем аргументы командной строки для необязательного указания другого файла.
import sys

# Подключаем чтение файлов Excel формата xlsx.
from openpyxl import load_workbook


# ПУНКТ 1. ФУНКЦИИ ФОРМИРОВАНИЯ МАТРИЦ РАССТОЯНИЙ

# 1.1. Манхэттенское расстояние равно сумме модулей разностей двух координат.
def manhattan_matrix(observations):
    # Определяем количество наблюдений и создаём квадратную матрицу нулей.
    n = len(observations)
    matrix = [[0 for _ in range(n)] for _ in range(n)]

    # Перебираем пары выше диагонали: каждую пару достаточно рассчитать один раз.
    for i in range(n):
        for j in range(i + 1, n):
            # Складываем абсолютные разности по двум координатам.
            distance = sum(
                abs(observations[i][k] - observations[j][k]) for k in range(2)
            )

            # Записываем расстояние в две симметричные ячейки матрицы.
            matrix[i][j] = distance
            matrix[j][i] = distance

    # Возвращаем готовую матрицу; её диагональ остаётся нулевой.
    return matrix


# 1.2. Расстояние Хэмминга равно числу несовпадающих координат: 0, 1 или 2.
def hamming_matrix(observations):
    # Определяем количество наблюдений и создаём квадратную матрицу нулей.
    n = len(observations)
    matrix = [[0 for _ in range(n)] for _ in range(n)]

    # Перебираем пары наблюдений выше главной диагонали.
    for i in range(n):
        for j in range(i + 1, n):
            # Считаем несовпадения: True даёт 1, False даёт 0 при суммировании.
            distance = sum(
                observations[i][k] != observations[j][k] for k in range(2)
            )

            # Заполняем симметричные ячейки; деление на число координат не нужно.
            matrix[i][j] = distance
            matrix[j][i] = distance

    # Возвращаем матрицу расстояний Хэмминга.
    return matrix


# 1.3. Евклидово расстояние равно корню из суммы квадратов разностей координат.
def euclidean_matrix(observations):
    # Определяем количество наблюдений и создаём квадратную матрицу нулей.
    n = len(observations)
    matrix = [[0 for _ in range(n)] for _ in range(n)]

    # Перебираем пары наблюдений выше главной диагонали.
    for i in range(n):
        for j in range(i + 1, n):
            # Вычисляем сумму квадратов разностей по двум координатам.
            squared_distance = sum(
                (observations[i][k] - observations[j][k]) ** 2
                for k in range(2)
            )

            # Извлекаем корень и записываем результат в симметричные ячейки.
            distance = sqrt(squared_distance)
            matrix[i][j] = distance
            matrix[j][i] = distance

    # Возвращаем матрицу без округления вычисленных значений.
    return matrix


# ПУНКТ 2. ВОСЕМЬ НАБЛЮДЕНИЙ В EXCEL И ЧТЕНИЕ ДАННЫХ ИЗ ФАЙЛА
# Файл lab1.xlsx подготовлен отдельно и включён в комплект работы.
# Лист Наблюдения: A2:A9 — метки, B2:C9 — две целочисленные координаты.

# Выполняем чтение и проверку при запуске файла как самостоятельной программы.
if __name__ == "__main__":
    # Используем указанный путь либо файл lab1.xlsx рядом с программой.
    excel_path = (Path(sys.argv[1]) if len(sys.argv) > 1
                  else Path(__file__).with_name("lab1.xlsx"))

    # Открываем существующий Excel-файл только для чтения.
    workbook = load_workbook(excel_path, read_only=True, data_only=True)

    # Подготавливаем пустые списки, которые будут заполнены данными из Excel.
    labels = []
    observations = []

    # Используем try/finally, чтобы закрыть файл даже при ошибке в данных.
    try:
        # Выбираем лист, содержащий восемь исходных наблюдений.
        sheet = workbook["Наблюдения"]

        # Считываем строки 2–9 и столбцы A–C: метку и две координаты.
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_row=9, min_col=1, max_col=3,
                            values_only=True), start=2
        ):
            # Разделяем считанную строку на метку, первую и вторую координату.
            label, x1, x2 = row

            # Проверяем наличие метки и двух конечных целочисленных значений.
            if label is None or str(label).strip() == "":
                raise ValueError(f"В строке {row_number} отсутствует метка.")
            if any(type(x) not in (int, float) or not isfinite(x)
                   or x != int(x) for x in (x1, x2)):
                raise ValueError(f"В строке {row_number} нужны две целые координаты.")

            # Добавляем наблюдение из Excel; метку храним отдельно от признаков.
            labels.append(str(label))
            observations.append([int(x1), int(x2)])
    finally:
        # Освобождаем открытый Excel-файл после чтения.
        workbook.close()

    # Проверяем уникальность меток, чтобы однозначно подписать строки матриц.
    if len(set(labels)) != 8:
        raise ValueError("Нужны восемь различных меток наблюдений.")

    # Показываем, какие именно наблюдения были считаны из Excel-файла.
    print("Работа 1. Матрица расстояний")
    print(f"Источник: {excel_path.name}, лист Наблюдения, диапазон A2:C9")
    print("Считаны 8 наблюдений с двумя координатами:")
    for label, point in zip(labels, observations):
        print(f"{label}: {point}")

    # ПУНКТ 3. ВЫЧИСЛЕНИЕ МАТРИЦ И ПРОВЕРКА ФУНКЦИЙ

    # Передаём всем трём функциям наблюдения, считанные на предыдущем шаге.
    matrices = {
        "Манхэттен": manhattan_matrix(observations),
        "Хэмминг": hamming_matrix(observations),
        "Евклид": euclidean_matrix(observations),
    }

    # Выводим матрицы в порядке задания; округляем только печать расстояний.
    for name, matrix in matrices.items():
        # Для евклидовой метрики показываем четыре знака после запятой.
        decimals = 4 if name == "Евклид" else 0

        # Печатаем название метрики, заголовки столбцов и строки матрицы.
        print(f"\n{name}")
        print(" " * 6 + "".join(f"{label:>9}" for label in labels))
        for label, row in zip(labels, matrix):
            print(f"{label:>6}" + "".join(f"{x:9.{decimals}f}" for x in row))

    # Задаём независимо рассчитанную контрольную матрицу Манхэттена.
    # Это ожидаемые РАССТОЯНИЯ для приложенного файла, а не исходные наблюдения.
    expected_manhattan = [
        [0, 20, 23, 15, 15, 12, 13, 30],
        [20, 0, 19, 23, 11, 16, 7, 12],
        [23, 19, 0, 18, 8, 35, 18, 15],
        [15, 23, 18, 0, 18, 27, 16, 33],
        [15, 11, 8, 18, 0, 27, 10, 15],
        [12, 16, 35, 27, 27, 0, 17, 28],
        [13, 7, 18, 16, 10, 17, 0, 17],
        [30, 12, 15, 33, 15, 28, 17, 0],
    ]

    # Задаём контрольную матрицу Хэмминга с числом различающихся координат.
    expected_hamming = [
        [0, 2, 2, 2, 1, 1, 2, 2],
        [2, 0, 2, 2, 2, 2, 2, 2],
        [2, 2, 0, 2, 2, 2, 2, 2],
        [2, 2, 2, 0, 2, 2, 2, 2],
        [1, 2, 2, 2, 0, 2, 2, 2],
        [1, 2, 2, 2, 2, 0, 2, 2],
        [2, 2, 2, 2, 2, 2, 0, 2],
        [2, 2, 2, 2, 2, 2, 2, 0],
    ]

    # Задаём целые квадраты евклидовых расстояний для удобной ручной проверки.
    expected_euclidean_squared = [
        [0, 208, 377, 117, 225, 144, 97, 578],
        [208, 0, 193, 325, 73, 160, 25, 122],
        [377, 193, 0, 194, 32, 617, 164, 137],
        [117, 325, 194, 0, 162, 477, 178, 545],
        [225, 73, 32, 162, 0, 369, 52, 113],
        [144, 160, 617, 477, 369, 0, 145, 554],
        [97, 25, 164, 178, 52, 145, 0, 205],
        [578, 122, 137, 545, 113, 554, 205, 0],
    ]

    # Объединяем эталоны и получаем евклидовы расстояния извлечением корней.
    expected = {
        "Манхэттен": expected_manhattan,
        "Хэмминг": expected_hamming,
        "Евклид": [[sqrt(x) for x in row] for row in expected_euclidean_squared],
    }

    # Начинаем проверку всех элементов каждой из трёх матриц.
    print("\nПроверки")
    for name, matrix in matrices.items():
        # Проверяем, что матрица содержит восемь строк и восемь столбцов.
        if len(matrix) != 8 or any(len(row) != 8 for row in matrix):
            raise AssertionError(f"{name}: неверный размер матрицы.")

        # Проверяем нули на диагонали, симметрию и неотрицательность расстояний.
        for i in range(8):
            if matrix[i][i] != 0:
                raise AssertionError(f"{name}: диагональ должна быть нулевой.")
            for j in range(8):
                if matrix[i][j] < 0 or matrix[i][j] != matrix[j][i]:
                    raise AssertionError(f"{name}: нарушены свойства матрицы.")

                # Сравниваем целые расстояния точно, евклидовы — с допуском 10⁻¹².
                actual = matrix[i][j]
                control = expected[name][i][j]
                equal = (isclose(actual, control, rel_tol=0, abs_tol=1e-12)
                         if name == "Евклид" else actual == control)

                # При расхождении указываем метрику и пару наблюдений.
                if not equal:
                    raise AssertionError(
                        f"{name}, {labels[i]}–{labels[j]}: получено {actual}, "
                        f"ожидалось {control}. Эталон относится к приложенному "
                        "набору; при изменении данных его нужно пересчитать."
                    )

        # Сообщаем об успешной проверке всех 64 элементов текущей матрицы.
        print(f"{name}: 64 элемента совпали с эталоном; свойства матрицы — OK")

    # Выводим итог только после успешного завершения всех проверок.
    print("Проверены 192 элемента. Все проверки пройдены.")
